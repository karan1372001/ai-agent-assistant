# =============================================================================
#  KARAN'S AI ASSISTANT - BACKEND (main.py)
# =============================================================================

from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional
import requests
import sqlite3
import re
import json
import math
import time
import os
import webbrowser
import subprocess
import uuid
import base64
import io
import smtplib
from email.mime.text import MIMEText
from dotenv import load_dotenv
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
from ddgs import DDGS
import pywhatkit
import dateparser
from dateparser.search import search_dates
from pypdf import PdfReader
from docx import Document

load_dotenv()

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

DB_PATH = "memory.db"


def db_execute(sql, params=(), fetch=None, commit=False):
    local_conn = sqlite3.connect(DB_PATH, check_same_thread=False, timeout=10)
    local_conn.execute("PRAGMA journal_mode=WAL")
    local_cursor = local_conn.cursor()
    local_cursor.execute(sql, params)
    result = None
    if fetch == "one":
        result = local_cursor.fetchone()
    elif fetch == "all":
        result = local_cursor.fetchall()
    if commit:
        local_conn.commit()
    local_conn.close()
    return result


_startup_conn = sqlite3.connect(DB_PATH, check_same_thread=False)
_startup_cursor = _startup_conn.cursor()

_startup_cursor.execute("""
CREATE TABLE IF NOT EXISTS history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    role TEXT,
    content TEXT,
    timestamp TEXT,
    embedding TEXT
)
""")
_startup_conn.commit()

for column_name, column_type in [("timestamp", "TEXT"), ("embedding", "TEXT"), ("session_id", "TEXT")]:
    try:
        _startup_cursor.execute(f"ALTER TABLE history ADD COLUMN {column_name} {column_type}")
        _startup_conn.commit()
    except Exception:
        pass

_startup_cursor.execute("""
CREATE TABLE IF NOT EXISTS facts (
    key TEXT PRIMARY KEY,
    value TEXT
)
""")
_startup_conn.commit()

_startup_cursor.execute("""
CREATE TABLE IF NOT EXISTS trusted_actions (
    tool_name TEXT,
    tool_arg_key TEXT,
    approval_count INTEGER,
    PRIMARY KEY (tool_name, tool_arg_key)
)
""")
_startup_conn.commit()

_startup_cursor.execute("""
CREATE TABLE IF NOT EXISTS reminders (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    message TEXT,
    remind_at TEXT,
    created_at TEXT,
    session_id TEXT,
    notified INTEGER DEFAULT 0
)
""")
_startup_conn.commit()
_startup_conn.close()


class ChatRequest(BaseModel):
    message: str
    session_id: Optional[str] = None

class ImageChatRequest(BaseModel):
    message: str
    image_base64: str
    session_id: Optional[str] = None

class DocumentChatRequest(BaseModel):
    message: str
    filename: str
    file_base64: str
    session_id: Optional[str] = None

class ApprovalRequest(BaseModel):
    action_id: str
    approved: bool

class DeleteReminderRequest(BaseModel):
    id: int


PENDING_ACTIONS = {}

TRUSTABLE_TOOLS = ["open_app", "open_website", "open_folder"]
APPROVAL_THRESHOLD = 3


def normalize_action_key(tool_arg):
    return tool_arg.strip().lower()


def get_approval_count(tool_name, tool_arg):
    key = normalize_action_key(tool_arg)
    row = db_execute(
        "SELECT approval_count FROM trusted_actions WHERE tool_name = ? AND tool_arg_key = ?",
        (tool_name, key), fetch="one"
    )
    return row[0] if row else 0


def increment_approval_count(tool_name, tool_arg):
    key = normalize_action_key(tool_arg)
    row = db_execute(
        "SELECT approval_count FROM trusted_actions WHERE tool_name = ? AND tool_arg_key = ?",
        (tool_name, key), fetch="one"
    )
    if row:
        db_execute(
            "UPDATE trusted_actions SET approval_count = approval_count + 1 WHERE tool_name = ? AND tool_arg_key = ?",
            (tool_name, key), commit=True
        )
    else:
        db_execute(
            "INSERT INTO trusted_actions (tool_name, tool_arg_key, approval_count) VALUES (?, ?, 1)",
            (tool_name, key), commit=True
        )


def is_trusted(tool_name, tool_arg):
    if tool_name not in TRUSTABLE_TOOLS:
        return False
    return get_approval_count(tool_name, tool_arg) >= APPROVAL_THRESHOLD


def save_fact(key, value):
    db_execute("INSERT OR REPLACE INTO facts (key, value) VALUES (?, ?)", (key, value), commit=True)


def get_all_facts():
    return db_execute("SELECT key, value FROM facts", fetch="all") or []


def detect_and_save_facts(message):
    message_lower = message.lower()

    lang_match = re.search(r"favou?rite (?:programming |coding )?language is (\w+)", message_lower)
    if lang_match:
        save_fact("favorite_language", lang_match.group(1).capitalize())

    name_match = re.search(r"my name is (\w+)", message_lower)
    if name_match:
        save_fact("name", name_match.group(1).capitalize())

    fullname_match = re.search(r"my full name is ([\w\s]+)", message_lower)
    if fullname_match:
        save_fact("full_name", fullname_match.group(1).strip().title())

    nickname_match = re.search(r"(?:call me|my nickname is) (\w+)", message_lower)
    if nickname_match:
        save_fact("nickname", nickname_match.group(1).upper())

    boss_match = re.search(r"\bi am your boss\b|\byou work for me\b", message_lower)
    if boss_match:
        save_fact("relationship", "The user is my boss.")

    assistant_name_match = re.search(
        r"(?:from now on[, ]+)?your name is (\w+)|i(?:'ll| will) call you (\w+)|call yourself (\w+)",
        message_lower
    )
    if assistant_name_match:
        new_name = next(g for g in assistant_name_match.groups() if g)
        save_fact("assistant_name", new_name.capitalize())


TIMEZONE_MAP = {
    "uk": "Europe/London", "united kingdom": "Europe/London", "england": "Europe/London",
    "london": "Europe/London", "india": "Asia/Kolkata", "usa": "America/New_York",
    "us": "America/New_York", "united states": "America/New_York", "new york": "America/New_York",
    "california": "America/Los_Angeles", "los angeles": "America/Los_Angeles",
    "japan": "Asia/Tokyo", "tokyo": "Asia/Tokyo", "australia": "Australia/Sydney",
    "sydney": "Australia/Sydney", "dubai": "Asia/Dubai", "uae": "Asia/Dubai",
    "germany": "Europe/Berlin", "france": "Europe/Paris", "canada": "America/Toronto",
    "singapore": "Asia/Singapore",
}


def get_time(location=None):
    if location:
        location_key = location.strip().lower()
        tz_name = TIMEZONE_MAP.get(location_key)
        if tz_name:
            now = datetime.now(ZoneInfo(tz_name))
            return now.strftime("%I:%M %p on %B %d, %Y") + f" ({location.strip()} time)"
        else:
            return f"I don't have timezone data for '{location.strip()}' yet, so I can't give an exact local time for it."
    return datetime.now().strftime("%I:%M %p on %B %d, %Y") + " (your local time)"


def calculate(expression):
    try:
        if re.match(r'^[0-9+\-*/(). ]+$', expression):
            result = eval(expression)
            return str(result)
        else:
            return "Invalid expression"
    except Exception:
        return "Error calculating that"


def search_web(query):
    max_attempts = 3
    for attempt in range(max_attempts):
        try:
            with DDGS() as ddgs:
                results = ddgs.text(query, max_results=3)
            if results:
                output = []
                for r in results:
                    title = r.get("title", "")
                    body = r.get("body", "")
                    output.append(f"{title}: {body}")
                return "\n".join(output)
            time.sleep(2)
        except Exception:
            time.sleep(2)
    return "I searched but couldn't retrieve results right now. Please try again in a moment."


WEATHER_CODES = {
    0: "clear sky", 1: "mainly clear", 2: "partly cloudy", 3: "overcast",
    45: "fog", 48: "depositing rime fog",
    51: "light drizzle", 53: "moderate drizzle", 55: "dense drizzle",
    61: "slight rain", 63: "moderate rain", 65: "heavy rain",
    71: "slight snow fall", 73: "moderate snow fall", 75: "heavy snow fall",
    80: "slight rain showers", 81: "moderate rain showers", 82: "violent rain showers",
    95: "thunderstorm", 96: "thunderstorm with slight hail", 99: "thunderstorm with heavy hail",
}


def get_weather(location):
    try:
        geo_response = requests.get(
            "https://geocoding-api.open-meteo.com/v1/search",
            params={"name": location, "count": 1}
        )
        geo_data = geo_response.json()

        if "results" not in geo_data or not geo_data["results"]:
            return f"I couldn't find a place called '{location}'. Could you check the spelling or try a nearby bigger city/town?"

        place = geo_data["results"][0]
        lat = place["latitude"]
        lon = place["longitude"]
        place_name = place.get("name", location)
        country = place.get("country", "")

        weather_response = requests.get(
            "https://api.open-meteo.com/v1/forecast",
            params={
                "latitude": lat,
                "longitude": lon,
                "current": "temperature_2m,relative_humidity_2m,weather_code,wind_speed_10m,apparent_temperature",
                "timezone": "auto",
            }
        )
        weather_data = weather_response.json()
        current = weather_data.get("current", {})

        temp = current.get("temperature_2m")
        feels_like = current.get("apparent_temperature")
        humidity = current.get("relative_humidity_2m")
        wind = current.get("wind_speed_10m")
        code = current.get("weather_code")
        condition = WEATHER_CODES.get(code, "unknown conditions")

        location_label = f"{place_name}, {country}" if country else place_name

        return (
            f"Current weather in {location_label}: {condition}, "
            f"temperature {temp}\u00b0C (feels like {feels_like}\u00b0C), "
            f"humidity {humidity}%, wind speed {wind} km/h."
        )
    except Exception as e:
        return f"Failed to get weather: {e}"


def normalize_time_text(text):
    text = re.sub(r'(\d{1,2})\.(\d{2})\s*(am|pm)', r'\1:\2 \3', text, flags=re.IGNORECASE)
    text = re.sub(r'\bafter\s+(\d+)', r'in \1', text, flags=re.IGNORECASE)
    return text


def parse_clock_time_directly(text):
    match = re.search(r'\b(\d{1,2}):?(\d{2})?\s*(am|pm)\b', text, re.IGNORECASE)
    if not match:
        return None, None

    hour = int(match.group(1))
    minute = int(match.group(2)) if match.group(2) else 0
    meridiem = match.group(3).lower()

    if hour == 12:
        hour = 0
    if meridiem == "pm":
        hour += 12

    now = datetime.now()
    candidate = now.replace(hour=hour, minute=minute, second=0, microsecond=0)
    if candidate <= now:
        candidate += timedelta(days=1)

    return candidate, match.group(0)


def clean_reminder_message(text, matched_text):
    message = text.replace(matched_text, "")
    message = re.sub(r'\bat\s*$', '', message, flags=re.IGNORECASE)
    message = re.sub(r'^(remind me( to)?|to remind me( to)?)\s+', '', message, flags=re.IGNORECASE).strip()
    message = re.sub(r'^(to|that|about|for|me to)\s+', '', message, flags=re.IGNORECASE).strip()
    message = re.sub(r'\s+(to|that|about|at|in)$', '', message, flags=re.IGNORECASE).strip()
    message = re.sub(r'\s+in\s+next\s*$', '', message, flags=re.IGNORECASE).strip()
    message = re.sub(r'\s+in\s+\d+\s*(min|mins|minute|minutes|hour|hours)?\s*$', '', message, flags=re.IGNORECASE).strip()
    message = re.sub(r'^(i am going to sleep,?\s*|wake me( up)?\s*)', '', message, flags=re.IGNORECASE).strip()
    message = message.strip(" ,.-\"'")
    if not message:
        message = "Wake up / reminder"
    return message


def parse_reminder(description):
    cleaned = normalize_time_text(description)

    direct_time, matched_text = parse_clock_time_directly(cleaned)
    if direct_time:
        message = clean_reminder_message(cleaned, matched_text)
        return direct_time, message

    results = search_dates(
        cleaned,
        settings={"PREFER_DATES_FROM": "future", "RETURN_AS_TIMEZONE_AWARE": False}
    )
    if not results:
        return None, description

    matched_text, remind_at = results[0]
    message = clean_reminder_message(cleaned, matched_text)
    return remind_at, message


def set_reminder(description, session_id=None):
    remind_at, message = parse_reminder(description)
    if not remind_at:
        return ("I couldn't figure out WHEN to remind you - try something like "
                "'remind me at 5pm to call mom' or 'remind me in 30 minutes to check the oven'.")

    created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    remind_at_str = remind_at.strftime("%Y-%m-%d %H:%M:%S")

    db_execute(
        "INSERT INTO reminders (message, remind_at, created_at, session_id, notified) VALUES (?, ?, ?, ?, 0)",
        (message, remind_at_str, created_at, session_id), commit=True
    )

    friendly_time = remind_at.strftime("%I:%M %p on %B %d, %Y")
    return f'Got it - I\'ll remind you to "{message}" at {friendly_time}.'


COMMON_APPS = {
    "vscode": "code",
    "vs code": "code",
    "visual studio code": "code",
    "chrome": "chrome",
    "google chrome": "chrome",
    "edge": "msedge",
    "microsoft edge": "msedge",
    "word": "winword",
    "microsoft word": "winword",
    "excel": "excel",
    "microsoft excel": "excel",
    "powerpoint": "powerpnt",
    "microsoft powerpoint": "powerpnt",
    "notepad": "notepad",
    "calculator": "calc",
    "paint": "mspaint",
    "explorer": "explorer",
    "file explorer": "explorer",
    "spotify": "spotify",
}


def open_app(app_name):
    try:
        key = app_name.strip().lower()
        actual_command = COMMON_APPS.get(key, app_name)
        os.system(f"start {actual_command}")
        return f"Opened {app_name}"
    except Exception as e:
        return f"Failed to open {app_name}: {e}"


COMMON_SITES = {
    "wikipedia": "wikipedia.org", "youtube": "youtube.com", "google": "google.com",
    "github": "github.com", "amazon": "amazon.com", "reddit": "reddit.com",
    "twitter": "twitter.com", "x": "x.com", "facebook": "facebook.com",
    "instagram": "instagram.com", "linkedin": "linkedin.com", "netflix": "netflix.com",
    "gmail": "gmail.com",
}


def open_website(url):
    try:
        url_clean = url.strip()
        if not url_clean.startswith("http"):
            if "." not in url_clean:
                site_key = url_clean.lower()
                domain = COMMON_SITES.get(site_key, f"{url_clean}.com")
                url_clean = "https://" + domain
            else:
                url_clean = "https://" + url_clean
        webbrowser.open(url_clean)
        return f"Opened {url_clean}"
    except Exception as e:
        return f"Failed to open website: {e}"


def play_youtube(song_name):
    try:
        pywhatkit.playonyt(song_name)
        return f"Playing '{song_name}' on YouTube"
    except Exception as e:
        return f"Failed to play song: {e}"

def open_folder(path):
    try:
        os.startfile(path)
        return f"Opened folder: {path}"
    except Exception as e:
        return f"Failed to open folder: {e}"


def determine_file_extension(description):
    description_lower = description.lower()
    if "c language" in description_lower or " c " in description_lower or description_lower.startswith("c "):
        return ".c"
    elif "python" in description_lower:
        return ".py"
    elif "javascript" in description_lower or "js" in description_lower:
        return ".js"
    elif "java" in description_lower:
        return ".java"
    elif "html" in description_lower:
        return ".html"
    else:
        return ".txt"


def determine_opener_app(description):
    description_lower = description.lower()
    if "vs code" in description_lower or "vscode" in description_lower or "visual studio code" in description_lower:
        return "vscode"
    elif "notepad" in description_lower:
        return "notepad"
    else:
        return "notepad"


def extract_custom_path(description):
    path_match = re.search(r'[a-zA-Z]:\\(?:[^\\/:*?"<>|\r\n]+\\?)+', description)
    if path_match:
        return path_match.group(0).rstrip("\\")
    return None


def extract_custom_filename(description):
    match = re.search(
        r'(?:save(?: it| this)?\s+as|name(?:d)?\s+it|call(?:ed)?\s+it)\s+'
        r'([a-zA-Z0-9_\-. ]+?)(?=\s+(?:in|at|on|to)\b|[,]|$)',
        description,
        re.IGNORECASE
    )
    if not match:
        return None

    captured = match.group(1).strip()
    captured = re.sub(r'\s+name$', '', captured, flags=re.IGNORECASE)
    return captured.replace(" ", "_")


def create_and_open_file(content, description):
    try:
        custom_filename = extract_custom_filename(description)

        if custom_filename:
            has_extension = re.search(r'\.[a-zA-Z0-9]{1,10}$', custom_filename)
            if has_extension:
                filename = custom_filename
            else:
                extension = determine_file_extension(description)
                filename = f"{custom_filename}{extension}"
        else:
            extension = determine_file_extension(description)
            filename = f"generated_file{extension}"

        custom_path = extract_custom_path(description)
        if custom_path:
            folder = custom_path
            os.makedirs(folder, exist_ok=True)
        else:
            folder = os.path.join(os.path.expanduser("~"), "Desktop")

        filepath = os.path.join(folder, filename)

        with open(filepath, "w") as f:
            f.write(content)

        opener = determine_opener_app(description)

        if opener == "vscode":
            subprocess.Popen(["code", filepath], shell=True)
            app_used = "VS Code"
        else:
            subprocess.Popen(["notepad.exe", filepath])
            app_used = "Notepad"

        return f"Created {filename} in {folder} and opened it in {app_used}"
    except Exception as e:
        return f"Failed to create file: {e}"


# =============================================================================
# EMAIL SENDING (THE NEW FEATURE)
# Lets you say "email my supervisor that I'm sick" - the AI writes the
# actual email itself, shows it to you for approval, and only sends it
# after you click Approve. Uses your Gmail App Password from the .env file.
# =============================================================================

EMAIL_ADDRESS = os.environ.get("ASSISTANT_EMAIL_ADDRESS")
EMAIL_APP_PASSWORD = os.environ.get("ASSISTANT_EMAIL_APP_PASSWORD")


def parse_email_tool_arg(tool_arg):
    to_match = re.search(r'to\s*=\s*([^\s,]+@[^\s,]+)', tool_arg, re.IGNORECASE)
    about_match = re.search(r'about\s*=\s*(.+)$', tool_arg, re.IGNORECASE)
    to_address = to_match.group(1).strip() if to_match else None
    about = about_match.group(1).strip() if about_match else tool_arg.strip()
    return to_address, about


def compose_email_content(about):
    content_prompt = [
        {"role": "system", "content": (
            "Write a short, professional email based on the user's description. "
            "Respond in EXACTLY this format and nothing else:\n"
            "SUBJECT: <a short subject line>\n"
            "BODY:\n<the full email body text, written professionally>"
        )},
        {"role": "user", "content": about}
    ]
    raw = ask_ollama_chat(content_prompt)
    subject_match = re.search(r'SUBJECT:\s*(.+)', raw)
    body_match = re.search(r'BODY:\s*(.*)', raw, re.DOTALL)
    subject = subject_match.group(1).strip() if subject_match else "Message from your assistant"
    body = body_match.group(1).strip() if body_match else raw.strip()
    return subject, body


def send_email_now(to_address, subject, body):
    if not EMAIL_ADDRESS or not EMAIL_APP_PASSWORD:
        return ("Email sending isn't set up yet - ASSISTANT_EMAIL_ADDRESS and "
                "ASSISTANT_EMAIL_APP_PASSWORD need to be in your .env file first.")
    if not to_address:
        return "I don't have a valid recipient email address, so I couldn't send this."
    try:
        msg = MIMEText(body or "")
        msg["Subject"] = subject or "Message from your assistant"
        msg["From"] = EMAIL_ADDRESS
        msg["To"] = to_address

        with smtplib.SMTP("smtp.gmail.com", 587) as server:
            server.starttls()
            server.login(EMAIL_ADDRESS, EMAIL_APP_PASSWORD)
            server.send_message(msg)

        return f"Email sent to {to_address} with subject '{subject}'"
    except Exception as e:
        return f"Failed to send email: {e}"


COMPUTER_CONTROL_TOOLS = ["open_app", "open_website", "play_youtube", "open_folder", "create_file", "send_email"]
MAX_TOOL_ROUNDS = 4


def get_embedding(text):
    response = requests.post(
        "http://localhost:11434/api/embeddings",
        json={"model": "nomic-embed-text", "prompt": text}
    )
    return response.json()["embedding"]


def cosine_similarity(a, b):
    dot = sum(x * y for x, y in zip(a, b))
    mag_a = math.sqrt(sum(x * x for x in a))
    mag_b = math.sqrt(sum(x * x for x in b))
    if mag_a == 0 or mag_b == 0:
        return 0
    return dot / (mag_a * mag_b)


def find_relevant_memories(current_message, limit=5):
    current_embedding = get_embedding(current_message)
    all_rows = db_execute("SELECT role, content, embedding FROM history WHERE embedding IS NOT NULL", fetch="all") or []
    scored = []
    for role, content, embedding_json in all_rows:
        if not embedding_json:
            continue
        stored_embedding = json.loads(embedding_json)
        score = cosine_similarity(current_embedding, stored_embedding)
        scored.append((score, role, content))
    scored.sort(key=lambda x: x[0], reverse=True)
    return scored[:limit]


MAX_DOCUMENT_CHARS = 12000


def extract_text_from_pdf(file_bytes):
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return "\n".join(text_parts)
    except Exception as e:
        return f"[Error reading PDF: {e}]"


def extract_text_from_docx(file_bytes):
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)
    except Exception as e:
        return f"[Error reading Word document: {e}]"


def extract_document_text(filename, file_base64):
    file_bytes = base64.b64decode(file_base64)
    filename_lower = filename.lower()

    if filename_lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif filename_lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif filename_lower.endswith(".txt"):
        try:
            return file_bytes.decode("utf-8", errors="ignore")
        except Exception as e:
            return f"[Error reading text file: {e}]"
    else:
        return None


SYSTEM_PROMPT = """You are a helpful assistant with access to tools, long-term memory, and computer control.

RULE: To use a tool, respond with ONLY tool lines in this exact format, one per line, nothing else:
TOOL: tool_name(argument)

If the user's request needs MULTIPLE actions done in order, output MULTIPLE tool lines, one per line, 
in the order they should happen.

IMPORTANT: Do NOT use any tool for general knowledge questions, trivia, fun facts, jokes, advice, or
anything you can simply answer from what you already know. Only use a tool when the request truly
needs it - real-time info (time/weather/search), math, reminders, or an actual action on the computer.

Available tools:
- get_time() -> current date and time where YOU are. get_time(location) -> the accurate current
  time in another place, e.g. get_time(UK) or get_time(India). ALWAYS use this tool for any
  question about the current time or date, anywhere - NEVER use search_web for time/date questions.
- get_weather(location) -> the REAL current weather (temperature, condition, humidity, wind) for
  any city/town, e.g. get_weather(Middlesbrough) or get_weather(Mumbai). ALWAYS use this tool for
  any question about current weather or temperature - NEVER use search_web for weather.
- calculate(expression) -> basic math
- search_web(query) -> search the internet for CURRENT/real-time info you don't already know
  (breaking news, live scores, etc) - NOT for general knowledge, trivia, or facts.
- set_reminder(description) -> creates a reminder. The description MUST include both WHAT to
  remind about and WHEN, in natural language, e.g. set_reminder(call mom at 5pm) or
  set_reminder(check the oven in 30 minutes) or set_reminder(wake me up in 5 minutes).
- send_email(to=recipient_email, about=what the email should say) -> composes and sends a REAL
  email. ALWAYS include a valid email address after "to=". Describe what the email should
  communicate after "about=" in your own words - you do NOT need to write the actual email text,
  a separate step composes a professional email for you based on your description.
- open_app(app_name) -> opens an application
- open_website(url) -> opens a website - only when the user specifically asks to open/visit a site
- play_youtube(song_name) -> searches and plays a video on YouTube
- open_folder(path) -> opens an EXISTING folder to browse files
- create_file(description) -> writes code/content to a file and opens it in notepad or vscode.
  If the user gives a filename (e.g. "save it as KD.html" or "save it as Karan")
  or a folder path (e.g. "E:\\MyFolder"), include that exact wording in the
  description so it gets saved with the right name and in the right place. If the user wants it
  opened in VS Code, include "vscode" or "vs code" in the description too, in the SAME tool call
  as create_file - do not use a separate open_app step for this.

Examples:
User: tell me a fun fact
You: Did you know honey never spoils? Archaeologists have found 3000-year-old honey in Egyptian tombs that's still edible!

User: open chrome
You: TOOL: open_app(chrome)

User: what time is it in the UK
You: TOOL: get_time(UK)

User: what's the weather like in Middlesbrough
You: TOOL: get_weather(Middlesbrough)

User: remind me at 5pm to call mom
You: TOOL: set_reminder(call mom at 5pm)

User: email my supervisor john@company.com that I'm sick and can't come to work
You: TOOL: send_email(to=john@company.com, about=I am sick today and unable to come to work)

User: play gta 6 trailer on youtube
You: TOOL: play_youtube(gta 6 trailer)

User: write addition code in C and open in vscode
You: TOOL: create_file(addition code in C, open in vscode)

User: write html code for a landing page, save it as KD.html in E:\Projects
You: TOOL: create_file(html code for a landing page, save it as KD.html in E:\\Projects)

You will be given "Known facts about the user" - these are permanent, reliable facts (name, 
preferences, etc.) that you should ALWAYS remember and use correctly, no matter how far back 
they were mentioned. Never contradict or forget these facts.

You may also be given "Relevant memories" from past conversations for additional context.

For anything else, just respond normally as a helpful assistant.
"""


def ask_ollama_chat(messages):
    response = requests.post(
        "http://localhost:11434/api/chat",
        json={"model": "llama3.1", "messages": messages, "stream": False}
    )
    return response.json()["message"]["content"]


def stream_ollama_chat(messages):
    response = requests.post(
        "http://localhost:11434/api/chat",
        json={"model": "llama3.1", "messages": messages, "stream": True},
        stream=True
    )
    for line in response.iter_lines():
        if not line:
            continue
        try:
            chunk = json.loads(line)
        except Exception:
            continue
        piece = chunk.get("message", {}).get("content", "")
        if piece:
            yield piece
        if chunk.get("done"):
            break


def fake_stream_text(full_text, chunk_size=3, delay=0.02):
    words = full_text.split(" ")
    buffer = ""
    for i, word in enumerate(words):
        buffer += word + (" " if i < len(words) - 1 else "")
        if (i + 1) % chunk_size == 0:
            yield buffer
            buffer = ""
            time.sleep(delay)
    if buffer:
        yield buffer


def save_message(role, content, session_id=None):
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    try:
        embedding = get_embedding(content)
        embedding_json = json.dumps(embedding)
    except Exception:
        embedding_json = None
    db_execute(
        "INSERT INTO history (role, content, timestamp, embedding, session_id) VALUES (?, ?, ?, ?, ?)",
        (role, content, timestamp, embedding_json, session_id), commit=True
    )


def create_permission_request(steps, session_id=None):
    first_tool_name, first_tool_arg = steps[0]
    remaining_steps = steps[1:]

    content_for_file = None
    email_to = None
    email_subject = None
    email_body = None

    if first_tool_name == "create_file":
        content_prompt = [
            {"role": "system", "content": f"Write ONLY the raw code/content for: {first_tool_arg}. No explanation, no markdown, no backticks."},
            {"role": "user", "content": first_tool_arg}
        ]
        content_for_file = ask_ollama_chat(content_prompt)

    elif first_tool_name == "send_email":
        email_to, about = parse_email_tool_arg(first_tool_arg)
        if email_to:
            email_subject, email_body = compose_email_content(about)

    action_id = str(uuid.uuid4())
    PENDING_ACTIONS[action_id] = {
        "tool_name": first_tool_name,
        "tool_arg": first_tool_arg,
        "content": content_for_file,
        "email_to": email_to,
        "email_subject": email_subject,
        "email_body": email_body,
        "remaining_steps": remaining_steps,
        "session_id": session_id,
    }

    description = f"{first_tool_name}({first_tool_arg})"
    step_info = f" (step 1 of {len(steps)})" if len(steps) > 1 else ""

    if first_tool_name == "create_file":
        preview = content_for_file[:150] + ("..." if len(content_for_file) > 150 else "")
        reply = f"I'd like to create a file with this content{step_info}:\n\n{preview}\n\nDo you approve?"
    elif first_tool_name == "send_email":
        if not email_to:
            reply = "I couldn't figure out who to send this to - please include their email address, e.g. 'email john@example.com that I'm sick'."
        else:
            reply = (
                f"I'd like to send this email{step_info}:\n\n"
                f"To: {email_to}\n"
                f"Subject: {email_subject}\n\n"
                f"{email_body}\n\n"
                f"Do you approve?"
            )
    else:
        reply = f"I'd like to: {description.replace('_', ' ')}{step_info}. Do you approve this action?"

    save_message("assistant", f"[Requested permission for: {description}]", session_id)

    return {
        "needs_permission": True,
        "action_id": action_id,
        "description": description,
        "reply": reply
    }


def run_single_tool(tool_name, tool_arg, content=None, email_to=None, email_subject=None, email_body=None):
    if tool_name == "open_app":
        return open_app(tool_arg)
    elif tool_name == "open_website":
        return open_website(tool_arg)
    elif tool_name == "play_youtube":
        return play_youtube(tool_arg)
    elif tool_name == "open_folder":
        return open_folder(tool_arg)
    elif tool_name == "create_file":
        return create_and_open_file(content, tool_arg)
    elif tool_name == "send_email":
        return send_email_now(email_to, email_subject, email_body)
    else:
        return "Unknown tool"


def run_tool_detection_round(ai_reply, messages, session_id=None):
    for round_number in range(MAX_TOOL_ROUNDS):
        tool_matches = re.findall(r"TOOLS?:\s*(\w+)\((.*)\)", ai_reply)

        if not tool_matches:
            return {"final_reply": ai_reply, "permission": None, "messages": messages}

        computer_steps = []
        tool_results = []

        for tool_name, tool_arg in tool_matches:
            if tool_name in COMPUTER_CONTROL_TOOLS:
                computer_steps.append((tool_name, tool_arg))
            elif tool_name == "get_time":
                tool_results.append(get_time(tool_arg.strip() if tool_arg.strip() else None))
            elif tool_name == "get_weather":
                tool_results.append(get_weather(tool_arg.strip()))
            elif tool_name == "calculate":
                tool_results.append(calculate(tool_arg))
            elif tool_name == "search_web":
                tool_results.append(search_web(tool_arg))
            elif tool_name == "set_reminder":
                tool_results.append(set_reminder(tool_arg.strip(), session_id))

        if computer_steps:
            auto_run_results = []
            remaining_computer_steps = list(computer_steps)

            while remaining_computer_steps:
                next_tool_name, next_tool_arg = remaining_computer_steps[0]
                if is_trusted(next_tool_name, next_tool_arg):
                    result = run_single_tool(next_tool_name, next_tool_arg)
                    auto_run_results.append(f"{next_tool_name}({next_tool_arg}): {result}")
                    save_message("assistant", f"Auto-approved (trusted action): {result}", session_id)
                    remaining_computer_steps.pop(0)
                else:
                    break

            if remaining_computer_steps:
                permission_data = create_permission_request(remaining_computer_steps, session_id)
                if auto_run_results:
                    permission_data["reply"] = "\n".join(auto_run_results) + "\n\n" + permission_data["reply"]
                return {"final_reply": None, "permission": permission_data, "messages": messages}
            else:
                final_text = "\n".join(auto_run_results)
                save_message("assistant", final_text, session_id)
                return {"final_reply": final_text, "permission": None, "messages": messages}

        if tool_results:
            messages.append({"role": "assistant", "content": ai_reply})
            messages.append({"role": "user", "content": f"Tool results: {'; '.join(tool_results)}. Now reply to me naturally using these results."})
            ai_reply = ask_ollama_chat(messages)
        else:
            return {"final_reply": ai_reply, "permission": None, "messages": messages}

    cleaned_reply = re.sub(r"TOOLS?:\s*\w+\(.*?\)\s*", "", ai_reply).strip()
    if not cleaned_reply:
        cleaned_reply = "I tried a few different approaches but couldn't get a clean final answer - could you try rephrasing your question?"
    return {"final_reply": cleaned_reply, "permission": None, "messages": messages}


def build_chat_messages(user_message, session_id=None):
    if session_id:
        recent = db_execute(
            "SELECT role, content FROM history WHERE session_id = ? ORDER BY id DESC LIMIT 6",
            (session_id,), fetch="all"
        ) or []
    else:
        recent = db_execute("SELECT role, content FROM history ORDER BY id DESC LIMIT 6", fetch="all") or []
    recent = list(recent)
    recent.reverse()

    relevant = find_relevant_memories(user_message, limit=5)

    facts = get_all_facts()
    assistant_name = None
    user_facts = []
    for key, value in facts:
        if key == "assistant_name":
            assistant_name = value
        else:
            user_facts.append((key, value))

    identity_prefix = ""
    if assistant_name:
        identity_prefix = (
            f"Your name is {assistant_name}. Whenever asked your name, or referring to "
            f"yourself, ALWAYS use {assistant_name} - this is permanent and never changes "
            f"unless the user explicitly renames you again.\n\n"
        )

    memory_context = ""
    if user_facts:
        memory_context += "Known facts about the user (ALWAYS remember these, never forget):\n"
        for key, value in user_facts:
            memory_context += f"- {key.replace('_', ' ')}: {value}\n"

    if relevant:
        memory_context += "\nRelevant memories from past conversations:\n"
        for score, role, content in relevant:
            memory_context += f"- {role}: {content}\n"

    messages = [{"role": "system", "content": identity_prefix + SYSTEM_PROMPT + "\n" + memory_context}]
    for role, content in recent:
        messages.append({"role": "user" if role == "user" else "assistant", "content": content})
    return messages


def try_direct_reminder(message):
    if re.search(r"\b(remind me|wake me( up)?|don'?t let me forget)\b", message, re.IGNORECASE):
        return set_reminder(message)
    return None


@app.get("/")
def read_root():
    return {"status": "Backend is running"}


@app.post("/chat")
def chat(request: ChatRequest):
    save_message("user", request.message, request.session_id)
    detect_and_save_facts(request.message)

    play_match = re.search(r"play (.+?) on youtube", request.message, re.IGNORECASE)
    if play_match:
        song_name = play_match.group(1)
        return create_permission_request([("play_youtube", song_name)], request.session_id)

    reminder_reply = try_direct_reminder(request.message)
    if reminder_reply:
        save_message("assistant", reminder_reply, request.session_id)
        return {"reply": reminder_reply, "needs_permission": False}

    messages = build_chat_messages(request.message, request.session_id)
    ai_reply = ask_ollama_chat(messages)

    result = run_tool_detection_round(ai_reply, messages, request.session_id)
    if result["permission"]:
        return result["permission"]

    final_reply = result["final_reply"]
    save_message("assistant", final_reply, request.session_id)
    return {"reply": final_reply, "needs_permission": False}


@app.post("/chat-stream")
def chat_stream(request: ChatRequest):
    session_id = request.session_id
    save_message("user", request.message, session_id)
    detect_and_save_facts(request.message)

    play_match = re.search(r"play (.+?) on youtube", request.message, re.IGNORECASE)
    if play_match:
        song_name = play_match.group(1)
        permission_data = create_permission_request([("play_youtube", song_name)], session_id)

        def permission_only_stream():
            yield "PERMISSION_JSON:" + json.dumps(permission_data)

        return StreamingResponse(permission_only_stream(), media_type="text/plain")

    reminder_reply = try_direct_reminder(request.message)
    if reminder_reply:
        def reminder_stream():
            for piece in fake_stream_text(reminder_reply):
                yield piece
            save_message("assistant", reminder_reply, session_id)

        return StreamingResponse(reminder_stream(), media_type="text/plain")

    messages = build_chat_messages(request.message, session_id)
    ai_reply = ask_ollama_chat(messages)

    result = run_tool_detection_round(ai_reply, messages, session_id)

    if result["permission"]:
        permission_data = result["permission"]

        def permission_only_stream():
            yield "PERMISSION_JSON:" + json.dumps(permission_data)

        return StreamingResponse(permission_only_stream(), media_type="text/plain")

    final_reply = result["final_reply"]
    updated_messages = result["messages"]
    tools_were_used = len(updated_messages) > len(messages)

    if tools_were_used:
        def real_stream():
            full_text = ""
            for piece in stream_ollama_chat(updated_messages):
                full_text += piece
                yield piece
            save_message("assistant", full_text, session_id)

        return StreamingResponse(real_stream(), media_type="text/plain")
    else:
        def replay_stream():
            for piece in fake_stream_text(final_reply):
                yield piece
            save_message("assistant", final_reply, session_id)

        return StreamingResponse(replay_stream(), media_type="text/plain")


@app.post("/approve-action")
def approve_action(request: ApprovalRequest):
    action = PENDING_ACTIONS.get(request.action_id)

    if not action:
        return {"reply": "This action has expired or doesn't exist anymore."}

    session_id = action.get("session_id")

    if not request.approved:
        del PENDING_ACTIONS[request.action_id]
        save_message("assistant", "Action was denied by the user. Remaining steps cancelled.", session_id)
        return {"reply": "Okay, I won't do that. I've also cancelled any remaining steps."}

    tool_name = action["tool_name"]
    tool_arg = action["tool_arg"]
    content = action.get("content")

    tool_result = run_single_tool(
        tool_name, tool_arg, content,
        email_to=action.get("email_to"),
        email_subject=action.get("email_subject"),
        email_body=action.get("email_body"),
    )

    remaining_steps = action.get("remaining_steps", [])
    del PENDING_ACTIONS[request.action_id]

    if tool_name in TRUSTABLE_TOOLS:
        increment_approval_count(tool_name, tool_arg)

    save_message("assistant", f"Action approved and completed: {tool_result}", session_id)

    if remaining_steps:
        next_request = create_permission_request(remaining_steps, session_id)
        next_request["reply"] = f"{tool_result}\n\nNext: {next_request['reply']}"
        return next_request

    return {"reply": tool_result, "needs_permission": False}


@app.post("/chat-image")
def chat_image(request: ImageChatRequest):
    save_message("user", request.message + " [sent an image]", request.session_id)
    response = requests.post(
        "http://localhost:11434/api/generate",
        json={"model": "llava", "prompt": request.message, "images": [request.image_base64], "stream": False}
    )
    result = response.json()
    if "response" not in result:
        return {"reply": f"Error from vision model: {result}"}
    ai_reply = result["response"]
    save_message("assistant", ai_reply, request.session_id)
    return {"reply": ai_reply}


@app.post("/chat-document")
def chat_document(request: DocumentChatRequest):
    save_message("user", request.message + f" [sent a document: {request.filename}]", request.session_id)

    document_text = extract_document_text(request.filename, request.file_base64)

    if document_text is None:
        reply = f"Sorry, I can only read .pdf, .docx, and .txt files right now - '{request.filename}' isn't one of those."
        save_message("assistant", reply, request.session_id)
        return {"reply": reply}

    if not document_text.strip():
        reply = f"I opened '{request.filename}' but couldn't find any readable text inside it - it might be a scanned image or an empty file."
        save_message("assistant", reply, request.session_id)
        return {"reply": reply}

    truncated = False
    if len(document_text) > MAX_DOCUMENT_CHARS:
        document_text = document_text[:MAX_DOCUMENT_CHARS]
        truncated = True

    truncation_note = "\n\n[Note: this document was long, so only the first portion is shown above.]" if truncated else ""

    document_prompt = [
        {"role": "system", "content": (
            "You are a helpful assistant. The user has shared a document with you. "
            "Answer their question using ONLY the document content below as your source of truth. "
            "If the answer isn't in the document, say so honestly instead of guessing."
        )},
        {"role": "user", "content": (
            f"Document filename: {request.filename}\n\n"
            f"Document content:\n{document_text}{truncation_note}\n\n"
            f"My question: {request.message}"
        )}
    ]

    ai_reply = ask_ollama_chat(document_prompt)
    save_message("assistant", ai_reply, request.session_id)
    return {"reply": ai_reply}


@app.get("/history")
def get_history():
    rows = db_execute("SELECT role, content, timestamp FROM history ORDER BY id ASC", fetch="all") or []
    return {"history": [{"role": r[0], "content": r[1], "timestamp": r[2]} for r in rows]}


@app.get("/trusted-actions")
def get_trusted_actions():
    rows = db_execute(
        "SELECT tool_name, tool_arg_key, approval_count FROM trusted_actions ORDER BY approval_count DESC",
        fetch="all"
    ) or []
    return {
        "trusted_actions": [
            {"tool_name": r[0], "action": r[1], "approval_count": r[2], "is_trusted": r[2] >= APPROVAL_THRESHOLD}
            for r in rows
        ],
        "threshold": APPROVAL_THRESHOLD,
    }


@app.get("/reminders")
def get_reminders():
    rows = db_execute(
        "SELECT id, message, remind_at, created_at, notified FROM reminders ORDER BY remind_at ASC",
        fetch="all"
    ) or []
    return {
        "reminders": [
            {"id": r[0], "message": r[1], "remind_at": r[2], "created_at": r[3], "notified": bool(r[4])}
            for r in rows
        ]
    }


@app.get("/reminders/due")
def get_due_reminders():
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    rows = db_execute(
        "SELECT id, message, remind_at FROM reminders WHERE remind_at <= ? AND notified = 0",
        (now_str,), fetch="all"
    ) or []
    due = [{"id": r[0], "message": r[1], "remind_at": r[2]} for r in rows]
    if due:
        ids = [r["id"] for r in due]
        placeholders = ",".join("?" * len(ids))
        db_execute(f"UPDATE reminders SET notified = 1 WHERE id IN ({placeholders})", tuple(ids), commit=True)
    return {"due": due}


@app.post("/reminders/delete")
def delete_reminder(request: DeleteReminderRequest):
    db_execute("DELETE FROM reminders WHERE id = ?", (request.id,), commit=True)
    return {"success": True}